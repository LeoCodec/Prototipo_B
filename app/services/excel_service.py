from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def build_excel(study, output_path, summary):
    wb = Workbook()
    ws = wb.active
    ws.title = "Estudio"

    rows = [
        ("Folio", study.folio),
        ("Alumno", study.student_name),
        ("Estado", study.status),
        ("Ingreso total", summary["total_income"]),
        ("Gasto total", summary["total_expenses"]),
        ("Disponible", summary["available"]),
        ("Integrantes", summary["household_size"]),
        ("Ingreso per cápita", summary["income_per_capita"]),
        ("Disponible per cápita", summary["available_per_capita"]),
        ("Carga de gasto (%)", summary["expense_load"]),
        ("Valoración final (manual)", study.final_fee if study.final_fee is not None else "Pendiente"),
    ]
    for r, (label, value) in enumerate(rows, start=1):
        ws.cell(r, 1, label)
        ws.cell(r, 2, value)
        ws.cell(r, 1).font = Font(bold=True)

    pages = wb.create_sheet("Paginas")
    pages.append(["Página", "Archivo", "Calidad", "Desenfoque", "Brillo", "OCR", "Texto confirmado", "Error OCR"])
    for image in study.images:
        pages.append([
            image.page_number, image.filename, image.quality_status,
            image.blur_score, image.brightness_score,
            image.ocr_text, image.confirmed_text, image.ocr_error
        ])

    qa = wb.create_sheet("Control de calidad")
    qa.append(["Regla", "Resultado"])
    qa.append(["Páginas capturadas", len(study.images)])
    qa.append(["Páginas ROJO", sum(1 for x in study.images if x.quality_status == "ROJO")])
    qa.append(["Páginas AMARILLO", sum(1 for x in study.images if x.quality_status == "AMARILLO")])
    qa.append(["Páginas VERDE", sum(1 for x in study.images if x.quality_status == "VERDE")])

    green = PatternFill("solid", fgColor="70AD47")
    for cell in ws[1]:
        cell.fill = green

    wb.save(output_path)
