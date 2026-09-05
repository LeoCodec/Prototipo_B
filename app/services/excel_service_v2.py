import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.socioeconomic_metrics import calculate_metrics


GREEN = "5F9700"
DARK_GREEN = "315500"
WHITE = "FFFFFF"


def _style_header(cell):
    cell.fill = PatternFill("solid", fgColor=DARK_GREEN)
    cell.font = Font(color=WHITE, bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _auto_width(ws, max_width=45):
    for col_cells in ws.columns:
        length = 0
        col_idx = col_cells[0].column
        for cell in col_cells:
            value = "" if cell.value is None else str(cell.value)
            length = max(length, len(value))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(length + 2, 10), max_width)


def _load_fields(instance_path, study_id):
    path = Path(instance_path) / "field_extractions" / f"study_{study_id}.json"
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {}


def build_workbook(study, instance_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"

    ws["A1"] = "IPPLIAP - Estudio Socioeconómico - Prototipo B"
    ws["A1"].font = Font(size=16, bold=True, color=DARK_GREEN)
    ws.merge_cells("A1:D1")

    rows = [
        ("Folio", study.folio),
        ("Identificador", study.student_name),
        ("Estatus", study.status),
        ("Ingreso total", study.total_income or 0),
        ("Gasto total", study.total_expenses or 0),
        ("Integrantes", study.household_size or 1),
        ("Cuota final autorizada", study.final_fee if study.final_fee is not None else ""),
    ]

    for idx, (label, value) in enumerate(rows, start=3):
        ws.cell(idx, 1, label)
        ws.cell(idx, 2, value)
        ws.cell(idx, 1).font = Font(bold=True, color=DARK_GREEN)

    metrics = calculate_metrics(
        study.total_income or 0,
        study.total_expenses or 0,
        study.household_size or 1,
    )

    metric_rows = [
        ("Disponible", metrics.available),
        ("Ingreso per cápita", metrics.income_per_capita),
        ("Disponible per cápita", metrics.available_per_capita),
        ("Carga de gasto (%)", metrics.expense_burden_pct),
    ]

    for idx, (label, value) in enumerate(metric_rows, start=3):
        ws.cell(idx, 3, label)
        ws.cell(idx, 4, value)
        ws.cell(idx, 3).font = Font(bold=True, color=DARK_GREEN)

    chart_start = 12
    ws.cell(chart_start, 1, "Indicador")
    ws.cell(chart_start, 2, "Monto")
    _style_header(ws.cell(chart_start, 1))
    _style_header(ws.cell(chart_start, 2))

    chart_data = [
        ("Ingreso", metrics.total_income),
        ("Gasto", metrics.total_expenses),
        ("Disponible", metrics.available),
    ]
    for r, (label, value) in enumerate(chart_data, start=chart_start + 1):
        ws.cell(r, 1, label)
        ws.cell(r, 2, value)

    chart = BarChart()
    chart.title = "Resumen económico"
    data = Reference(ws, min_col=2, min_row=chart_start, max_row=chart_start + len(chart_data))
    cats = Reference(ws, min_col=1, min_row=chart_start + 1, max_row=chart_start + len(chart_data))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, "F12")

    pages = wb.create_sheet("Paginas")
    for c, header in enumerate(["Página","Archivo","Calidad","Desenfoque","Brillo","Error OCR"], start=1):
        pages.cell(1, c, header)
        _style_header(pages.cell(1, c))

    for r, image in enumerate(study.images, start=2):
        vals = [
            image.page_number,
            image.filename,
            image.quality_status,
            image.blur_score,
            image.brightness_score,
            image.ocr_error or "",
        ]
        for c, value in enumerate(vals, start=1):
            pages.cell(r, c, value)

    ocr = wb.create_sheet("OCR")
    for c, header in enumerate(["Página","Texto OCR","Texto confirmado"], start=1):
        ocr.cell(1, c, header)
        _style_header(ocr.cell(1, c))

    for r, image in enumerate(study.images, start=2):
        ocr.cell(r, 1, image.page_number)
        ocr.cell(r, 2, image.ocr_text or "")
        ocr.cell(r, 3, image.confirmed_text or "")
        ocr.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ocr.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top")

    fields_sheet = wb.create_sheet("Campos_confirmados")
    payload = _load_fields(instance_path, study.id)
    for c, header in enumerate(["Campo","Detectado","Confirmado","Confianza"], start=1):
        fields_sheet.cell(1, c, header)
        _style_header(fields_sheet.cell(1, c))

    for r, (name, item) in enumerate(payload.get("fields", {}).items(), start=2):
        fields_sheet.cell(r, 1, name)
        fields_sheet.cell(r, 2, item.get("detected", ""))
        fields_sheet.cell(r, 3, item.get("confirmed", ""))
        fields_sheet.cell(r, 4, item.get("confidence", 0))

    ind = wb.create_sheet("Indicadores")
    for c, header in enumerate(["Indicador","Valor","Fórmula"], start=1):
        ind.cell(1, c, header)
        _style_header(ind.cell(1, c))

    indicators = [
        ("Ingreso total", metrics.total_income, "Suma de ingresos"),
        ("Gasto total", metrics.total_expenses, "Suma de gastos"),
        ("Disponible", metrics.available, "Ingreso - Gasto"),
        ("Ingreso per cápita", metrics.income_per_capita, "Ingreso / integrantes"),
        ("Disponible per cápita", metrics.available_per_capita, "Disponible / integrantes"),
        ("Carga de gasto (%)", metrics.expense_burden_pct, "Gasto / Ingreso * 100"),
    ]

    for r, row in enumerate(indicators, start=2):
        for c, value in enumerate(row, start=1):
            ind.cell(r, c, value)

    ws["A20"] = "Nota:"
    ws["B20"] = "Los indicadores son descriptivos. La cuota final permanece bajo decisión humana."

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        _auto_width(sheet)

    return wb
