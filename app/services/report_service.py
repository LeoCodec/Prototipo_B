from __future__ import annotations

import json
import re
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.socioeconomic_metrics import (
    calculate_metrics,
    descriptive_findings,
)


DARK = "315500"
WHITE = "FFFFFF"


def safe_name(value: str) -> str:
    value = re.sub(r'[\\/:*?"<>|]+', "_", value or "caso")
    value = re.sub(r"\s+", "_", value).strip("._")
    return value[:80] or "caso"


def field_payload(instance_path: str, study_id: int) -> dict:
    path = (
        Path(instance_path)
        / "field_extractions"
        / f"study_{study_id}.json"
    )

    if not path.exists():
        return {
            "fields": {},
            "context_notes": "",
            "review_notes": "",
            "analysis": {},
        }

    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {
            "fields": {},
            "context_notes": "",
            "review_notes": "",
            "analysis": {},
        }


def _metric_values(study, payload):
    analysis = payload.get("analysis", {})

    income = analysis.get(
        "total_income",
        getattr(study, "total_income", 0) or 0,
    )
    expenses = analysis.get(
        "total_expenses",
        getattr(study, "total_expenses", 0) or 0,
    )
    household = analysis.get(
        "household_size",
        getattr(study, "household_size", 1) or 1,
    )
    final_fee = analysis.get(
        "final_fee",
        getattr(study, "final_fee", None),
    )

    return (
        calculate_metrics(income, expenses, household),
        final_fee,
    )


def _header(cell):
    cell.fill = PatternFill("solid", fgColor=DARK)
    cell.font = Font(color=WHITE, bold=True)
    cell.alignment = Alignment(horizontal="center")


def _autowidth(ws, max_width=55):
    for cells in ws.columns:
        index = cells[0].column
        size = max(len(str(cell.value or "")) for cell in cells)
        ws.column_dimensions[get_column_letter(index)].width = min(
            max(size + 2, 10),
            max_width,
        )


def case_workbook(study, instance_path: str) -> BytesIO:
    payload = field_payload(instance_path, study.id)
    metrics, final_fee = _metric_values(study, payload)

    wb = Workbook()

    # 1. Resumen
    ws = wb.active
    ws.title = "Resumen"

    ws["A1"] = "IPPLIAP - Prototipo B - Resumen del expediente"
    ws["A1"].font = Font(size=16, bold=True, color=DARK)
    ws.merge_cells("A1:D1")

    rows = [
        ("Folio", study.folio),
        ("Caso", study.student_name),
        ("Estatus", study.status),
        ("Ingreso total", metrics.total_income),
        ("Gasto total", metrics.total_expenses),
        ("Disponible", metrics.available),
        ("Integrantes del hogar", metrics.household_size),
        ("Ingreso per capita", metrics.income_per_capita),
        ("Disponible per capita", metrics.available_per_capita),
        ("Carga de gasto (%)", metrics.expense_burden_pct),
        ("Cuota final autorizada", final_fee if final_fee is not None else ""),
        ("Observaciones de contexto", payload.get("context_notes", "")),
        ("Notas de revision", payload.get("review_notes", "")),
    ]

    for row_index, (label, value) in enumerate(rows, start=3):
        ws.cell(row_index, 1, label).font = Font(
            bold=True,
            color=DARK,
        )
        ws.cell(row_index, 2, value)

    # Grafico simple y descriptivo.
    base = 18
    ws.cell(base, 1, "Indicador")
    ws.cell(base, 2, "Monto")
    _header(ws.cell(base, 1))
    _header(ws.cell(base, 2))

    chart_rows = [
        ("Ingreso", metrics.total_income),
        ("Gasto", metrics.total_expenses),
        ("Disponible", metrics.available),
    ]
    for row_index, values in enumerate(chart_rows, start=base + 1):
        ws.cell(row_index, 1, values[0])
        ws.cell(row_index, 2, values[1])

    chart = BarChart()
    chart.title = "Resumen economico"
    data = Reference(
        ws,
        min_col=2,
        min_row=base,
        max_row=base + len(chart_rows),
    )
    categories = Reference(
        ws,
        min_col=1,
        min_row=base + 1,
        max_row=base + len(chart_rows),
    )
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    ws.add_chart(chart, "D18")

    # 2. Campos confirmados
    fields_ws = wb.create_sheet("Campos_confirmados")
    for col, title in enumerate(
        ["Campo", "Detectado", "Confirmado", "Confianza"],
        start=1,
    ):
        fields_ws.cell(1, col, title)
        _header(fields_ws.cell(1, col))

    for row_index, (name, item) in enumerate(
        payload.get("fields", {}).items(),
        start=2,
    ):
        fields_ws.cell(
            row_index,
            1,
            item.get("label") or name,
        )
        fields_ws.cell(row_index, 2, item.get("detected", ""))
        fields_ws.cell(row_index, 3, item.get("confirmed", ""))
        fields_ws.cell(
            row_index,
            4,
            round(float(item.get("confidence", 0)) * 100, 2),
        )

    # 3. OCR por pagina
    ocr_ws = wb.create_sheet("OCR_por_pagina")
    for col, title in enumerate(
        ["Pagina", "Archivo", "Calidad", "Texto OCR", "Texto confirmado", "Error OCR"],
        start=1,
    ):
        ocr_ws.cell(1, col, title)
        _header(ocr_ws.cell(1, col))

    for row_index, image in enumerate(study.images, start=2):
        values = [
            image.page_number,
            image.filename,
            image.quality_status,
            image.ocr_text or "",
            image.confirmed_text or "",
            image.ocr_error or "",
        ]

        for col, value in enumerate(values, start=1):
            ocr_ws.cell(row_index, col, value)

        ocr_ws.cell(row_index, 4).alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )
        ocr_ws.cell(row_index, 5).alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )

    # 4. Indicadores
    ind_ws = wb.create_sheet("Indicadores")
    indicators = [
        ("Ingreso total", metrics.total_income, "IT"),
        ("Gasto total", metrics.total_expenses, "GT"),
        ("Disponible", metrics.available, "IT - GT"),
        ("Ingreso per capita", metrics.income_per_capita, "IT / N"),
        ("Disponible per capita", metrics.available_per_capita, "(IT - GT) / N"),
        ("Carga de gasto (%)", metrics.expense_burden_pct, "(GT / IT) x 100"),
    ]

    for col, title in enumerate(["Indicador", "Resultado", "Formula"], start=1):
        ind_ws.cell(1, col, title)
        _header(ind_ws.cell(1, col))

    for row_index, values in enumerate(indicators, start=2):
        for col, value in enumerate(values, start=1):
            ind_ws.cell(row_index, col, value)

    # 5. Hallazgos descriptivos
    findings_ws = wb.create_sheet("Lectura_descriptiva")
    findings_ws["A1"] = "Hallazgos descriptivos"
    _header(findings_ws["A1"])
    findings_ws["B1"] = "Nota"
    _header(findings_ws["B1"])

    for row_index, item in enumerate(
        descriptive_findings(metrics),
        start=2,
    ):
        findings_ws.cell(row_index, 1, row_index - 1)
        findings_ws.cell(row_index, 2, item)

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        _autowidth(sheet)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def bulk_workbook(studies, instance_path: str) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen_general"

    headers = [
        "Folio",
        "Caso",
        "Estatus",
        "Ingreso total",
        "Gasto total",
        "Disponible",
        "Integrantes",
        "Ingreso per capita",
        "Disponible per capita",
        "Carga de gasto (%)",
        "Cuota final autorizada",
        "Contexto humano",
    ]

    for col, title in enumerate(headers, start=1):
        ws.cell(1, col, title)
        _header(ws.cell(1, col))

    for row_index, study in enumerate(studies, start=2):
        payload = field_payload(instance_path, study.id)
        metrics, final_fee = _metric_values(study, payload)

        values = [
            study.folio,
            study.student_name,
            study.status,
            metrics.total_income,
            metrics.total_expenses,
            metrics.available,
            metrics.household_size,
            metrics.income_per_capita,
            metrics.available_per_capita,
            metrics.expense_burden_pct,
            final_fee if final_fee is not None else "",
            payload.get("context_notes", ""),
        ]

        for col, value in enumerate(values, start=1):
            ws.cell(row_index, col, value)

    _autowidth(ws)
    ws.freeze_panes = "A2"

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
