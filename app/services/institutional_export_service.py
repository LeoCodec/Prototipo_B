from __future__ import annotations

from io import BytesIO
from statistics import mean, median
import re

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.services.institutional_form_service import load_form

GREEN = "2E5D0B"
LIGHT_GREEN = "EAF4D8"
WHITE = "FFFFFF"

LABELS = {
    "solicitud_tipo": "Tipo de solicitud",
    "turno": "Turno",
    "grado": "Grado que solicita",
    "fecha": "Fecha",
    "nombre_alumno": "Nombre",
    "apellido_paterno_alumno": "Apellido paterno",
    "apellido_materno_alumno": "Apellido materno",
    "fecha_nacimiento": "Fecha de nacimiento",
    "edad": "Edad",
    "domicilio": "Domicilio",
    "telefono": "Teléfono",
    "celular": "Celular",
    "solicitante_nombre": "Nombre de quien solicita",
    "solicitante_parentesco": "Parentesco",
    "solicitante_edad": "Edad del solicitante",
    "solicitante_domicilio": "Domicilio del solicitante",
    "solicitante_telefono": "Teléfono",
    "solicitante_celular": "Celular",
    "solicitante_escolaridad": "Escolaridad",
    "solicitante_ocupacion": "Ocupación",
    "solicitante_trabajo": "Nombre y domicilio del trabajo",
    "solicitante_trabajo_telefono": "Teléfono de trabajo",
    "solicitante_correo": "Correo electrónico",
    "razon_apoyo": "Razón de solicitud",
    "vivienda_tenencia": "Tenencia de vivienda",
    "vivienda_tipo": "Tipo de vivienda",
    "habitaciones": "Habitaciones",
    "dormitorios": "Dormitorios",
    "espacios_vivienda": "Espacios",
    "paredes": "Paredes",
    "techos": "Techos",
    "pisos": "Pisos",
    "servicios_vivienda": "Servicios",
    "bienes_hogar": "Bienes del hogar",
    "zona": "Características de la zona",
    "comunidad": "Servicios/comunidad",
    "transporte": "Transporte",
    "transporte_modelo": "Modelo",
    "transporte_anio": "Año",
    "referencias": "Referencias",
    "beca_actual_pct": "Beca actual (%)",
    "predial_anual": "Predial anual",
    "gastos_anuales": "Otros gastos anuales",
    "vacaciones_anual": "Vacaciones anual",
    "renta": "Renta mensual",
    "luz": "Luz mensual",
    "agua": "Agua mensual",
    "telefono_gasto": "Teléfono mensual",
    "gas": "Gas mensual",
    "alimentos": "Alimentos mensual",
    "automovil": "Automóvil mensual",
    "pasajes": "Pasajes mensual",
    "colegiaturas": "Colegiaturas mensual",
    "vestido": "Vestido mensual",
    "medico_medicinas": "Médico y medicinas mensual",
    "muebles_hogar": "Muebles del hogar mensual",
    "creditos_personales": "Créditos personales mensual",
    "otros_gastos": "Otros gastos mensual",
    "servicio_medico": "Servicio médico",
    "enfermedades_cronicas": "Enfermedades crónicas",
    "tiempo_libre": "Tiempo libre",
    "observaciones": "Observaciones",
    "contexto_sociofamiliar": "Contexto sociofamiliar",
    "final_fee": "Cuota final autorizada",
}

MONTHLY_EXPENSE_LABELS = [
    ("renta", "Renta"),
    ("luz", "Luz"),
    ("agua", "Agua"),
    ("telefono_gasto", "Teléfono"),
    ("gas", "Gas"),
    ("alimentos", "Alimentos"),
    ("automovil", "Automóvil"),
    ("pasajes", "Pasajes"),
    ("colegiaturas", "Colegiaturas"),
    ("vestido", "Vestido"),
    ("medico_medicinas", "Médico y medicinas"),
    ("muebles_hogar", "Muebles del hogar"),
    ("creditos_personales", "Créditos personales"),
    ("otros_gastos", "Otros gastos"),
]

ANNUAL_EQ_LABELS = [
    ("predial_anual_mensual_equiv", "Predial (equiv. mensual)"),
    ("gastos_anuales_mensual_equiv", "Gastos anuales (equiv. mensual)"),
    ("vacaciones_anual_mensual_equiv", "Vacaciones (equiv. mensual)"),
]


def money(v):
    if v is None or v == "":
        return "Pendiente"
    try:
        return "$ {:,.2f} MXN".format(float(v))
    except Exception:
        return str(v)


def _num(v, default=0.0):
    if v is None or v == "":
        return default
    try:
        if isinstance(v, (int, float)):
            return float(v)
        t = str(v).strip().replace("$", "").replace("MXN", "").replace(" ", "")
        if "," in t and "." in t:
            t = t.replace(",", "") if t.rfind(".") > t.rfind(",") else t.replace(".", "").replace(",", ".")
        elif "," in t:
            parts = t.split(",")
            t = "".join(parts[:-1]) + "." + parts[-1] if len(parts[-1]) == 2 else t.replace(",", "")
        return float(t)
    except Exception:
        return default


def _safe(t):
    return (re.sub(r"[^A-Za-z0-9_-]+", "_", str(t or "").strip())[:28] or "Caso")


def _header(ws, row, cols):
    for c, v in enumerate(cols, 1):
        x = ws.cell(row=row, column=c, value=v)
        x.fill = PatternFill("solid", fgColor=GREEN)
        x.font = Font(color=WHITE, bold=True)
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _fit(ws):
    for c in range(1, ws.max_column + 1):
        w = 12
        for r in range(1, min(ws.max_row, 160) + 1):
            v = ws.cell(r, c).value
            if v is not None:
                w = max(w, min(len(str(v)) + 2, 44))
        ws.column_dimensions[get_column_letter(c)].width = w


def _money_cell(cell, value):
    cell.value = None if value is None else _num(value)
    if value is not None:
        cell.number_format = '$#,##0.00 "MXN"'


def _section(wb, title, pairs):
    ws = wb.create_sheet(title)
    _header(ws, 1, ["Campo", "Valor"])
    r = 2
    for a, b in pairs:
        ws.cell(r, 1, a)
        ws.cell(r, 2, b)
        r += 1
    ws.freeze_panes = "A2"
    _fit(ws)
    return ws


def _link(paragraph, text, url):
    rid = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    h.append(run)
    paragraph._p.append(h)


def _doc_table(doc, rows, headers=("Campo", "Valor")):
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = headers[0]
    t.rows[0].cells[1].text = headers[1]
    for a, b in rows:
        if b in (None, ""):
            continue
        c = t.add_row().cells
        c[0].text = str(a)
        c[1].text = str(b)
    return t


def payload(study, instance_path):
    p = load_form(instance_path, study.id)
    return p, p.get("data", {}), p.get("calculation", {})


def _expense_rows(c):
    detail = c.get("expense_detail") or {}
    total_expenses = _num(c.get("total_expenses"), 0.0)
    total_income = _num(c.get("total_income"), 0.0)
    rows = []

    for key, label in ANNUAL_EQ_LABELS + MONTHLY_EXPENSE_LABELS:
        value = max(_num(detail.get(key), 0.0), 0.0)
        rows.append(
            (
                label,
                value,
                (value / total_expenses * 100.0) if total_expenses > 0 else 0.0,
                (value / total_income * 100.0) if total_income > 0 else 0.0,
            )
        )
    return rows


def _income_rows(p, total_household_income):
    rows = []
    for row in p.get("income_rows") or []:
        name = str(row.get("integrante") or "").strip() or "Sin nombre"
        income = max(_num(row.get("ingreso_mensual"), 0.0), 0.0)
        contribution = _num(row.get("aportacion_mensual"), None)
        if contribution is None:
            contribution = income
        contribution = max(contribution, 0.0)
        rows.append(
            (
                name,
                income,
                contribution,
                (contribution / income * 100.0) if income > 0 else 0.0,
                (contribution / total_household_income * 100.0) if total_household_income > 0 else 0.0,
            )
        )
    return rows


def build_case_excel(study, instance_path):
    p, d, c = payload(study, instance_path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"

    ws["A1"] = "IPPLIAP - Análisis económico del estudio socioeconómico"
    ws["A1"].font = Font(size=16, bold=True, color=GREEN)
    ws.merge_cells("A1:B1")

    _header(ws, 3, ["Dato", "Valor"])
    metadata = [
        ("Folio", study.folio),
        ("Caso", study.student_name),
        ("Estatus", study.status),
        ("Turno", d.get("turno", "")),
        ("Fecha", d.get("fecha", "")),
    ]
    rr = 4
    for label, value in metadata:
        ws.cell(rr, 1, label)
        ws.cell(rr, 2, value)
        rr += 1

    _header(ws, 10, ["Indicador", "Resultado"])
    metrics = [
        ("Ingreso total", c.get("total_income", 0)),
        ("Gasto mensual equivalente", c.get("total_expenses", 0)),
        ("Disponible", c.get("available", 0)),
        ("Integrantes del hogar", c.get("household_size", 1)),
        ("Ingreso per cápita", c.get("income_per_capita", 0)),
        ("Disponible per cápita", c.get("available_per_capita", 0)),
        ("Carga de gasto (%)", c.get("expense_burden_pct", 0)),
        ("Cuota de referencia", c.get("reference_fee", 0)),
        ("Cuota final autorizada", c.get("final_fee")),
    ]
    money_labels = {
        "Ingreso total",
        "Gasto mensual equivalente",
        "Disponible",
        "Ingreso per cápita",
        "Disponible per cápita",
        "Cuota de referencia",
        "Cuota final autorizada",
    }
    rr = 11
    for label, value in metrics:
        ws.cell(rr, 1, label)
        if label in money_labels:
            _money_cell(ws.cell(rr, 2), value)
        else:
            ws.cell(rr, 2, value)
            if label == "Carga de gasto (%)":
                ws.cell(rr, 2).number_format = '0.00"%"'
        rr += 1

    ws["A22"] = "Lectura del cálculo"
    ws["A22"].font = Font(bold=True, color=GREEN, size=12)
    ws["A23"] = c.get("explanation", "")
    ws.merge_cells("A23:B25")
    ws["A23"].alignment = Alignment(wrap_text=True, vertical="top")

    maps = p.get("maps_url") or ""
    ws["A27"] = "Google Maps (referencia)"
    ws["B27"] = "Abrir ubicación"
    if maps:
        ws["B27"].hyperlink = maps
        ws["B27"].style = "Hyperlink"

    analysis = wb.create_sheet("Analisis_economico")
    _header(analysis, 1, ["Gasto", "Monto mensual", "% del gasto", "% del ingreso"])
    expense_rows = _expense_rows(c)
    r = 2
    for label, value, pct_expense, pct_income in expense_rows:
        analysis.cell(r, 1, label)
        _money_cell(analysis.cell(r, 2), value)
        analysis.cell(r, 3, pct_expense / 100.0)
        analysis.cell(r, 3).number_format = "0.00%"
        analysis.cell(r, 4, pct_income / 100.0)
        analysis.cell(r, 4).number_format = "0.00%"
        r += 1

    total_household_income = _num(c.get("total_income"), 0.0)
    _header(analysis, 1, ["Gasto", "Monto mensual", "% del gasto", "% del ingreso", "", "Integrante", "Ingreso mensual", "Aportación mensual", "% aportado de su ingreso", "% aportado al hogar"])
    income_rows = _income_rows(p, total_household_income)
    for idx, row in enumerate(income_rows, 2):
        name, income, contribution, pct_self, pct_home = row
        analysis.cell(idx, 6, name)
        _money_cell(analysis.cell(idx, 7), income)
        _money_cell(analysis.cell(idx, 8), contribution)
        analysis.cell(idx, 9, pct_self / 100.0)
        analysis.cell(idx, 9).number_format = "0.00%"
        analysis.cell(idx, 10, pct_home / 100.0)
        analysis.cell(idx, 10).number_format = "0.00%"

    # Top 5 gastos
    top_start = max(r, 2 + len(income_rows)) + 3
    _header(analysis, top_start, ["Cinco gastos principales", "Monto", "% del gasto"])
    for idx, row in enumerate(sorted(expense_rows, key=lambda x: x[1], reverse=True)[:5], top_start + 1):
        analysis.cell(idx, 1, row[0])
        _money_cell(analysis.cell(idx, 2), row[1])
        analysis.cell(idx, 3, row[2] / 100.0)
        analysis.cell(idx, 3).number_format = "0.00%"

    detail = wb.create_sheet("Detalle_calculo")
    _header(detail, 1, ["Parámetro", "Valor"])
    detail_rows = [
        ("Folio", study.folio),
        ("Caso", study.student_name),
        ("Turno", d.get("turno", "")),
        ("Integrantes usados en el cálculo", c.get("household_size", 1)),
        ("Ingreso total", c.get("total_income", 0)),
        ("Gasto mensual equivalente", c.get("total_expenses", 0)),
        ("Disponible", c.get("available", 0)),
        ("Ingreso per cápita", c.get("income_per_capita", 0)),
        ("Disponible per cápita", c.get("available_per_capita", 0)),
        ("Carga de gasto (%)", c.get("expense_burden_pct", 0)),
        ("Tope de cuota configurado", c.get("max_reference_fee", 0)),
        ("Porcentaje de ingreso aplicado", c.get("income_ratio", 0)),
        ("Porcentaje de disponibilidad aplicado", c.get("disposable_ratio", 0)),
        ("Beca actual (%)", c.get("scholarship_pct", 0)),
        ("Tope por beca", c.get("scholarship_cap", 0)),
        ("Cuota de referencia", c.get("reference_fee", 0)),
        ("Cuota final autorizada", c.get("final_fee")),
        ("Explicación", c.get("explanation", "")),
    ]
    money_detail = {
        "Ingreso total",
        "Gasto mensual equivalente",
        "Disponible",
        "Ingreso per cápita",
        "Disponible per cápita",
        "Tope de cuota configurado",
        "Tope por beca",
        "Cuota de referencia",
        "Cuota final autorizada",
    }
    rr = 2
    for label, value in detail_rows:
        detail.cell(rr, 1, label)
        if label in money_detail:
            _money_cell(detail.cell(rr, 2), value)
        else:
            detail.cell(rr, 2, value)
            if "Porcentaje" in label and value not in (None, ""):
                detail.cell(rr, 2).number_format = "0.00%"
        rr += 1

    # Charts
    if any(value > 0 for _, value, _, _ in expense_rows):
        chart = DoughnutChart()
        chart.title = "Distribución del gasto familiar"
        chart.style = 10
        chart.height = 8
        chart.width = 12
        data_ref = Reference(analysis, min_col=2, min_row=2, max_row=1 + len(expense_rows))
        cat_ref = Reference(analysis, min_col=1, min_row=2, max_row=1 + len(expense_rows))
        chart.add_data(data_ref, titles_from_data=False)
        chart.set_categories(cat_ref)
        chart.holeSize = 55
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
        ws.add_chart(chart, "D3")

    if income_rows:
        chart2 = BarChart()
        chart2.type = "col"
        chart2.style = 10
        chart2.title = "Ingreso y aportación por integrante"
        chart2.y_axis.title = "MXN"
        chart2.x_axis.title = "Integrante"
        data_ref = Reference(analysis, min_col=7, max_col=8, min_row=1, max_row=1 + len(income_rows))
        cat_ref = Reference(analysis, min_col=6, min_row=2, max_row=1 + len(income_rows))
        chart2.add_data(data_ref, titles_from_data=True)
        chart2.set_categories(cat_ref)
        chart2.height = 8
        chart2.width = 13
        ws.add_chart(chart2, "D19")

    chart3 = BarChart()
    chart3.type = "col"
    chart3.style = 10
    chart3.title = "Ingreso, gasto y disponibilidad"
    chart3.y_axis.title = "MXN"
    data_ref = Reference(ws, min_col=2, min_row=11, max_row=13)
    cat_ref = Reference(ws, min_col=1, min_row=11, max_row=13)
    chart3.add_data(data_ref, titles_from_data=False)
    chart3.set_categories(cat_ref)
    chart3.height = 8
    chart3.width = 12
    ws.add_chart(chart3, "D35")

    ws.freeze_panes = "A3"
    analysis.freeze_panes = "A2"
    detail.freeze_panes = "A2"
    _fit(ws)
    _fit(analysis)
    _fit(detail)

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def build_case_word(study, instance_path):
    p, d, c = payload(study, instance_path)
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("ESTUDIO SOCIOECONÓMICO")
    r.bold = True
    r.font.size = Pt(18)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Instituto Pedagógico para Problemas del Lenguaje, I.A.P.").bold = True

    doc.add_paragraph("Folio: " + str(study.folio))
    doc.add_paragraph("Caso: " + str(study.student_name))
    doc.add_paragraph("Estatus: " + str(study.status))

    doc.add_heading("1. Datos generales del alumno", 1)
    akeys = ["solicitud_tipo","turno","grado","fecha","nombre_alumno","apellido_paterno_alumno","apellido_materno_alumno","fecha_nacimiento","edad","domicilio","telefono","celular"]
    _doc_table(doc, [(LABELS[k], d.get(k, "")) for k in akeys])

    maps = p.get("maps_url") or ""
    if maps:
        par = doc.add_paragraph("Ubicación de referencia: ")
        _link(par, "Abrir en Google Maps", maps)
        doc.add_paragraph("La ubicación es aproximada y no confirma por sí sola la vivienda.")

    doc.add_heading("2. Datos de quien solicita la beca", 1)
    skeys = ["solicitante_nombre","solicitante_parentesco","solicitante_edad","solicitante_domicilio","solicitante_telefono","solicitante_celular","solicitante_escolaridad","solicitante_ocupacion","solicitante_trabajo","solicitante_trabajo_telefono","solicitante_correo","razon_apoyo"]
    _doc_table(doc, [(LABELS[k], d.get(k, "")) for k in skeys])

    doc.add_heading("3. Integrantes del hogar", 1)
    fam = p.get("family_members") or []
    if fam:
        t = doc.add_table(rows=1, cols=4)
        t.style = "Table Grid"
        for i, x in enumerate(("Nombre", "Parentesco", "Edad", "Ocupación")):
            t.rows[0].cells[i].text = x
        for row in fam:
            cc = t.add_row().cells
            for i, key in enumerate(("nombre", "parentesco", "edad", "ocupacion")):
                cc[i].text = str(row.get(key, ""))
    else:
        doc.add_paragraph("Sin integrantes registrados.")

    doc.add_heading("4. Vivienda", 1)
    vkeys = ["vivienda_tenencia","vivienda_tipo","habitaciones","dormitorios","espacios_vivienda","paredes","techos","pisos","servicios_vivienda","bienes_hogar","zona","comunidad","transporte","transporte_modelo","transporte_anio","referencias"]
    _doc_table(doc, [(LABELS[k], d.get(k, "")) for k in vkeys])

    doc.add_heading("5. Situación económica", 1)
    rows = p.get("income_rows") or []
    if rows:
        t = doc.add_table(rows=1, cols=6)
        t.style = "Table Grid"
        for i, x in enumerate(("Integrante", "Lugar de trabajo", "Antigüedad", "Puesto", "Ingreso", "Aportación")):
            t.rows[0].cells[i].text = x
        for row in rows:
            cc = t.add_row().cells
            vals = [
                row.get("integrante", ""),
                row.get("lugar_trabajo", ""),
                row.get("antiguedad", ""),
                row.get("puesto", ""),
                money(row.get("ingreso_mensual")),
                money(row.get("aportacion_mensual")),
            ]
            for i, v in enumerate(vals):
                cc[i].text = str(v)
    doc.add_paragraph("Ingreso mensual considerado: " + money(c.get("total_income", 0)))

    doc.add_heading("6. Distribución del gasto familiar", 1)
    ekeys = ["predial_anual","gastos_anuales","vacaciones_anual","renta","luz","agua","telefono_gasto","gas","alimentos","automovil","pasajes","colegiaturas","vestido","medico_medicinas","muebles_hogar","creditos_personales","otros_gastos"]
    _doc_table(doc, [(LABELS[k], money(d.get(k)) if d.get(k) not in (None, "") else "") for k in ekeys])

    doc.add_heading("7. Salud y contexto", 1)
    ckeys = ["servicio_medico","enfermedades_cronicas","tiempo_libre","observaciones","contexto_sociofamiliar","beca_actual_pct"]
    _doc_table(doc, [(LABELS[k], d.get(k, "")) for k in ckeys])

    doc.add_heading("8. Indicadores socioeconómicos", 1)
    _doc_table(
        doc,
        [
            ("Ingreso total", money(c.get("total_income", 0))),
            ("Gasto mensual equivalente", money(c.get("total_expenses", 0))),
            ("Disponible", money(c.get("available", 0))),
            ("Integrantes del hogar", c.get("household_size", 1)),
            ("Ingreso per cápita", money(c.get("income_per_capita", 0))),
            ("Disponible per cápita", money(c.get("available_per_capita", 0))),
            ("Carga de gasto", str(c.get("expense_burden_pct", 0)) + " %"),
        ],
        ("Indicador", "Resultado"),
    )

    doc.add_heading("9. Lectura económica y valoración", 1)
    doc.add_paragraph(
        "El hogar registra {n} integrante(s). El ingreso mensual considerado es {income}; "
        "el gasto mensual equivalente es {expenses}; y la disponibilidad posterior a gastos es {available}. "
        "El ingreso per cápita es {ipc} y la disponibilidad per cápita es {dpc}. "
        "La carga de gasto representa {burden:.2f} % del ingreso considerado.".format(
            n=c.get("household_size", 1),
            income=money(c.get("total_income", 0)),
            expenses=money(c.get("total_expenses", 0)),
            available=money(c.get("available", 0)),
            ipc=money(c.get("income_per_capita", 0)),
            dpc=money(c.get("available_per_capita", 0)),
            burden=_num(c.get("expense_burden_pct"), 0.0),
        )
    )
    _doc_table(
        doc,
        [
            ("Cuota de referencia calculada", money(c.get("reference_fee", 0))),
            ("Cuota final autorizada", money(c.get("final_fee"))),
        ],
        ("Concepto", "Resultado"),
    )
    doc.add_paragraph(str(c.get("explanation", "")))
    doc.add_paragraph(
        "La cuota de referencia es orientativa. La cuota final debe ser revisada y autorizada "
        "por el personal institucional responsable, considerando también las circunstancias familiares "
        "y la valoración profesional."
    )

    doc.add_heading("10. Alcance", 1)
    doc.add_paragraph(
        "Este documento resume la información capturada manualmente en el formulario digital. "
        "No sustituye la entrevista, la observación profesional ni la decisión institucional final."
    )

    out = BytesIO()
    doc.save(out)
    out.seek(0)
    return out


def _metric(vals):
    vals = [float(v) for v in vals if v is not None]
    return (0, None, None, None, None) if not vals else (len(vals), mean(vals), median(vals), min(vals), max(vals))


def build_bulk_excel(studies, instance_path):
    wb = Workbook()
    stats = wb.active
    stats.title = "Estadísticas"
    cases = wb.create_sheet("Casos")
    records = [(s, *payload(s, instance_path)) for s in studies]

    _header(stats, 1, ["Indicador", "Casos con dato", "Media", "Mediana", "Mínimo", "Máximo"])
    defs = [
        ("Ingreso total", "total_income", "money"),
        ("Gasto mensual equivalente", "total_expenses", "money"),
        ("Disponible", "available", "money"),
        ("Ingreso per cápita", "income_per_capita", "money"),
        ("Disponible per cápita", "available_per_capita", "money"),
        ("Carga de gasto (%)", "expense_burden_pct", "pct"),
        ("Cuota de referencia", "reference_fee", "money"),
        ("Cuota final autorizada", "final_fee", "money"),
    ]
    for i, (label, key, kind) in enumerate(defs, 2):
        count, avg, med, mn, mx = _metric([c.get(key) for _, _, _, c in records])
        stats.cell(i, 1, label)
        stats.cell(i, 2, count)
        for col, v in enumerate((avg, med, mn, mx), 3):
            stats.cell(i, col, v)
            if v is not None:
                stats.cell(i, col).number_format = '$#,##0.00 "MXN"' if kind == "money" else '0.00"%"'

    headers = [
        "Folio","Caso","Estatus","Turno","Ingreso total","Gasto mensual","Disponible",
        "Integrantes","Ingreso per cápita","Disponible per cápita","Carga de gasto (%)",
        "Cuota de referencia","Cuota final",
    ]
    _header(cases, 1, headers)
    for r, (s, p, d, c) in enumerate(records, 2):
        vals = [
            s.folio, s.student_name, s.status, d.get("turno", ""),
            c.get("total_income", 0), c.get("total_expenses", 0), c.get("available", 0),
            c.get("household_size", 1), c.get("income_per_capita", 0), c.get("available_per_capita", 0),
            c.get("expense_burden_pct", 0), c.get("reference_fee", 0), c.get("final_fee"),
        ]
        for col, v in enumerate(vals, 1):
            cases.cell(r, col, v)
        for col in (5, 6, 7, 9, 10, 12, 13):
            if cases.cell(r, col).value is not None:
                cases.cell(r, col).number_format = '$#,##0.00 "MXN"'
        cases.cell(r, 11).number_format = '0.00"%"'

    # Gráfica comparativa general
    if records:
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "Comparación por caso: ingreso, gasto y disponible"
        chart.y_axis.title = "MXN"
        chart.x_axis.title = "Caso"
        data_ref = Reference(cases, min_col=5, max_col=7, min_row=1, max_row=1 + len(records))
        cat_ref = Reference(cases, min_col=2, min_row=2, max_row=1 + len(records))
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cat_ref)
        chart.height = 9
        chart.width = 16
        stats.add_chart(chart, "H2")

        fee_chart = BarChart()
        fee_chart.type = "col"
        fee_chart.style = 10
        fee_chart.title = "Cuotas por caso"
        fee_chart.y_axis.title = "MXN"
        data_ref = Reference(cases, min_col=12, max_col=13, min_row=1, max_row=1 + len(records))
        fee_chart.add_data(data_ref, titles_from_data=True)
        fee_chart.set_categories(cat_ref)
        fee_chart.height = 9
        fee_chart.width = 16
        stats.add_chart(fee_chart, "H20")

    for idx, (s, p, d, c) in enumerate(records, 1):
        ws = wb.create_sheet(("Caso_%s_%s" % (idx, _safe(s.student_name)))[:31])
        ws["A1"] = f"{s.folio} - {s.student_name}"
        ws["A1"].font = Font(bold=True, size=14, color=GREEN)
        ws.merge_cells("A1:D1")

        _header(ws, 3, ["Indicador", "Resultado"])
        summary_rows = [
            ("Turno", d.get("turno", "")),
            ("Ingreso total", c.get("total_income", 0)),
            ("Gasto mensual", c.get("total_expenses", 0)),
            ("Disponible", c.get("available", 0)),
            ("Integrantes", c.get("household_size", 1)),
            ("Ingreso per cápita", c.get("income_per_capita", 0)),
            ("Disponible per cápita", c.get("available_per_capita", 0)),
            ("Carga de gasto (%)", c.get("expense_burden_pct", 0)),
            ("Cuota de referencia", c.get("reference_fee", 0)),
            ("Cuota final", c.get("final_fee")),
        ]
        rr = 4
        for label, value in summary_rows:
            ws.cell(rr, 1, label)
            ws.cell(rr, 2, value)
            if label in {"Ingreso total","Gasto mensual","Disponible","Ingreso per cápita","Disponible per cápita","Cuota de referencia","Cuota final"} and value is not None:
                ws.cell(rr, 2).number_format = '$#,##0.00 "MXN"'
            elif label == "Carga de gasto (%)":
                ws.cell(rr, 2).number_format = '0.00"%"'
            rr += 1

        start = rr + 2
        _header(ws, start, ["Gasto", "Monto mensual", "% del gasto"])
        exp_rows = _expense_rows(c)
        for ridx, (label, value, pct_exp, _) in enumerate(exp_rows, start + 1):
            ws.cell(ridx, 1, label)
            _money_cell(ws.cell(ridx, 2), value)
            ws.cell(ridx, 3, pct_exp / 100.0)
            ws.cell(ridx, 3).number_format = "0.00%"

        _fit(ws)

    stats.freeze_panes = "A2"
    cases.freeze_panes = "A2"
    _fit(stats)
    _fit(cases)

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out

