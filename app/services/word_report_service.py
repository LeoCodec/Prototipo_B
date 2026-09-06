from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.shared import Pt

from app.services.report_service import field_payload
from app.services.socioeconomic_metrics import (
    calculate_metrics,
    descriptive_findings,
)


def build_case_docx(study, instance_path: str) -> BytesIO:
    payload = field_payload(instance_path, study.id)
    analysis = payload.get("analysis", {})

    income = analysis.get(
        "total_income",
        getattr(study, "total_income", 0) or 0,
    )
    expenses = analysis.get(
        "total_expenses",
        getattr(study, "total_expenses", 0) or 0,
    )
    household = analysis.get(
        "household_size",
        getattr(study, "household_size", 1) or 1,
    )
    final_fee = analysis.get(
        "final_fee",
        getattr(study, "final_fee", None),
    )

    metrics = calculate_metrics(
        income,
        expenses,
        household,
    )

    doc = Document()
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10)

    doc.add_heading("IPPLIAP - Prototipo B", level=1)
    doc.add_paragraph(
        "Resumen de digitalizacion del estudio socioeconomico"
    )

    doc.add_heading("Identificacion del expediente", level=2)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"

    identification = [
        ("Folio", study.folio),
        ("Caso", study.student_name),
        ("Estatus", study.status),
    ]

    for label, value in identification:
        cells = table.add_row().cells
        cells[0].text = str(label)
        cells[1].text = str(value or "")

    doc.add_heading("Campos confirmados", level=2)
    confirmed_table = doc.add_table(rows=1, cols=3)
    confirmed_table.style = "Table Grid"
    confirmed_table.rows[0].cells[0].text = "Campo"
    confirmed_table.rows[0].cells[1].text = "Detectado"
    confirmed_table.rows[0].cells[2].text = "Confirmado"

    for name, item in payload.get("fields", {}).items():
        cells = confirmed_table.add_row().cells
        cells[0].text = str(item.get("label") or name)
        cells[1].text = str(item.get("detected") or "")
        cells[2].text = str(item.get("confirmed") or "")

    doc.add_heading("Indicadores descriptivos", level=2)
    metric_table = doc.add_table(rows=0, cols=2)
    metric_table.style = "Table Grid"

    metric_rows = [
        ("Ingreso total", f"${metrics.total_income:,.2f}"),
        ("Gasto total", f"${metrics.total_expenses:,.2f}"),
        ("Disponible", f"${metrics.available:,.2f}"),
        ("Integrantes del hogar", str(metrics.household_size)),
        ("Ingreso per capita", f"${metrics.income_per_capita:,.2f}"),
        ("Disponible per capita", f"${metrics.available_per_capita:,.2f}"),
        ("Carga de gasto", f"{metrics.expense_burden_pct:.2f}%"),
        (
            "Cuota final autorizada",
            "" if final_fee is None else f"${float(final_fee):,.2f}",
        ),
    ]

    for label, value in metric_rows:
        cells = metric_table.add_row().cells
        cells[0].text = label
        cells[1].text = value

    doc.add_heading("Lectura descriptiva", level=2)
    for finding in descriptive_findings(metrics):
        doc.add_paragraph(finding, style="List Bullet")

    doc.add_heading("Observaciones de contexto", level=2)
    doc.add_paragraph(
        payload.get("context_notes")
        or "Sin observaciones adicionales registradas."
    )

    doc.add_heading("Notas de revision del piloto", level=2)
    doc.add_paragraph(
        payload.get("review_notes")
        or "Sin notas adicionales de revision."
    )

    doc.add_heading("Alcance de la salida", level=2)
    doc.add_paragraph(
        "El OCR y los indicadores son herramientas de apoyo. "
        "La informacion manuscrita debe ser revisada por una persona "
        "y la cuota final no se determina automaticamente."
    )

    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output
